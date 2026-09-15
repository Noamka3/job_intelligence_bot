"""Raw text extraction from an uploaded resume file (PDF or DOCX)."""

from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader


class UnsupportedResumeFormatError(ValueError):
    def __init__(self, filename: str) -> None:
        super().__init__(f"Unsupported resume format: {filename!r} (expected .pdf or .docx)")


class UnreadableResumeError(ValueError):
    def __init__(self, filename: str, cause: Exception) -> None:
        super().__init__(
            f"Could not read {filename!r} as a {filename.rsplit('.', 1)[-1].upper()} file "
            f"(corrupt, encrypted, or mislabeled?): {type(cause).__name__}: {cause}"
        )


def extract_text(filename: str, content: bytes) -> str:
    suffix = filename.rsplit(".", maxsplit=1)[-1].lower() if "." in filename else ""
    if suffix == "pdf":
        extractor = _extract_pdf_text
    elif suffix == "docx":
        extractor = _extract_docx_text
    else:
        raise UnsupportedResumeFormatError(filename)

    try:
        return extractor(content)
    except Exception as exc:  # noqa: BLE001 - third-party parsers on untrusted bytes raise a zoo of types
        raise UnreadableResumeError(filename, exc) from exc


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _extract_docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)
