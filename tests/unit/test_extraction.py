from __future__ import annotations

import io

import pytest
from docx import Document

from app.services.candidate import extraction
from app.services.candidate.extraction import UnsupportedResumeFormatError, extract_text


def test_extract_text_rejects_unsupported_extension() -> None:
    with pytest.raises(UnsupportedResumeFormatError):
        extract_text("resume.txt", b"plain text")


def test_extract_text_rejects_missing_extension() -> None:
    with pytest.raises(UnsupportedResumeFormatError):
        extract_text("resume", b"plain text")


def test_extract_text_docx_roundtrip() -> None:
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Junior Software Engineer")
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("resume.docx", buffer.getvalue())

    assert "Jane Doe" in text
    assert "Junior Software Engineer" in text


def test_extract_text_pdf_joins_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    class FakePage:
        def __init__(self, text: str | None) -> None:
            self._text = text

        def extract_text(self) -> str | None:
            return self._text

    class FakeReader:
        def __init__(self, _stream: object) -> None:
            self.pages = [FakePage("Page one"), FakePage(None), FakePage("Page three")]

    monkeypatch.setattr(extraction, "PdfReader", FakeReader)

    text = extract_text("resume.pdf", b"%PDF-fake")

    assert text == "Page one\n\nPage three"
